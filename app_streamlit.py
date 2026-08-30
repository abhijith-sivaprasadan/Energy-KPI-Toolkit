#!/usr/bin/env python3
import datetime as dt
import hashlib
import json
import os
import tempfile
from io import BytesIO
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import streamlit as st
from openpyxl import load_workbook

from run_toolkit import (
    Params,
    compute_kpis,
    make_pdf_report,
    make_plots,
    parse_period_to_datetime,
    read_data_from_workbook,
    read_params,
    write_results_to_workbook,
)

REQUIRED_COLUMNS = {
    "Period",
    "Electricity_MWh",
    "NaturalGas_MWh",
    "Steam_MWh",
    "Production_t",
    "OperatingHours_h",
}

ACTION_COLUMNS = [
    "CreatedOn",
    "Period",
    "IssueType",
    "Severity",
    "SuggestedInvestigation",
    "Owner",
    "Status",
    "Notes",
]

PROFILE_DEFAULTS = {
    "Balanced (Recommended)": {
        "driver_model": "Both",
        "rolling_window": 3,
        "anomaly_z": 2.5,
        "drift_window": 3,
        "drift_threshold_z": 1.5,
    },
    "Conservative Alerts": {
        "driver_model": "Both",
        "rolling_window": 4,
        "anomaly_z": 3.0,
        "drift_window": 4,
        "drift_threshold_z": 2.0,
    },
    "Sensitive Alerts": {
        "driver_model": "Both",
        "rolling_window": 2,
        "anomaly_z": 2.0,
        "drift_window": 2,
        "drift_threshold_z": 1.0,
    },
}

REFERENCE_LINKS = [
    ("ISO 50001:2018 (EnMS framework)", "https://www.iso.org/standard/69426.html"),
    ("ISO 50006:2023 (EnPI/EnB guidance)", "https://www.iso.org/standard/79367.html"),
    (
        "U.S. DOE FEMP M&V Guidelines 5.0 (Options A/B/C/D)",
        "https://www.energy.gov/sites/default/files/2024-10/mv_guide_5_0.pdf",
    ),
    (
        "DOE: Measurement and Verification Options (Option C overview)",
        "https://www.energy.gov/femp/measurement-and-verification-options-federal-energy-and-water-saving-projects",
    ),
    (
        "E.S. Page (1954) CUSUM, Biometrika, DOI:10.1093/biomet/41.1-2.100",
        "https://academic.oup.com/biomet/article/41/1-2/100/456627",
    ),
    (
        "Kissock et al., ASHRAE RP-1050 inverse-modeling toolkit report",
        "https://hdl.handle.net/1969.1/2847",
    ),
    (
        "ORNL/DOE: regression-based tracking vs classic intensity (2018)",
        "https://www.osti.gov/pages/biblio/1474575",
    ),
    (
        "SSAB sustainability reports and energy-management directives",
        "https://www.ssab.com/en/company/sustainability/reports-and-documents",
    ),
]

PAPER_PLOT_MAPPING = [
    {
        "Diagnostic plot": "Actual vs Expected + R²/MAE",
        "Method concept": "Baseline-normalized regression EnPI fit quality",
        "Primary source": "ISO 50006:2023; ORNL/DOE regression tracking paper",
        "Why used here": "Quantifies how well selected drivers explain energy and whether model assumptions are reasonable.",
    },
    {
        "Diagnostic plot": "Residual control view (time-series with ±2σ/±3σ)",
        "Method concept": "Statistical process control on model residuals",
        "Primary source": "CUSUM/SPC tradition (Page, 1954); DOE M&V practice",
        "Why used here": "Separates random variation from potential special-cause events or operational shifts.",
    },
    {
        "Diagnostic plot": "Residual distribution histogram",
        "Method concept": "Residual diagnostics / model adequacy check",
        "Primary source": "Regression diagnostics guidance in M&V workflows",
        "Why used here": "Helps detect skewness/heavy tails that can bias alerting thresholds.",
    },
    {
        "Diagnostic plot": "Energy vs Production / Energy vs Hours",
        "Method concept": "Driver-response sensitivity and explanatory relationship",
        "Primary source": "ISO 50006 EnPI driver guidance; ASHRAE inverse modeling toolkit",
        "Why used here": "Shows whether selected operational drivers have plausible association with total energy.",
    },
    {
        "Diagnostic plot": "Utility share stacked area (electricity/gas/steam)",
        "Method concept": "Energy review and significant energy use breakdown",
        "Primary source": "ISO 50001 energy review principles",
        "Why used here": "Tracks structural shifts in source mix that can explain KPI movement.",
    },
    {
        "Diagnostic plot": "Rolling correlation (Energy-Production, Energy-Hours)",
        "Method concept": "Non-stationarity / regime change monitoring",
        "Primary source": "M&V ongoing tracking guidance; drift monitoring practice",
        "Why used here": "Highlights periods where historical driver relationships weaken or invert.",
    },
    {
        "Diagnostic plot": "Monthly seasonality boxplot",
        "Method concept": "Seasonal variation assessment",
        "Primary source": "Facility energy analytics practice under ISO 50006 frameworks",
        "Why used here": "Prevents false interpretation of seasonal patterns as efficiency regressions.",
    },
]


def _auth_gate() -> None:
    app_password = os.getenv("APP_PASSWORD", "").strip()
    if not app_password:
        return
    st.sidebar.header("Access")
    entered = st.sidebar.text_input("App password", type="password")
    if entered != app_password:
        st.warning("Authentication required.")
        st.stop()


def _load_csv(uploaded_file, demo_file: str = "demo_industrial_energy_data.csv") -> pd.DataFrame:
    if uploaded_file is None:
        return pd.read_csv(demo_file)
    return pd.read_csv(uploaded_file)


def _prepare_data(df: pd.DataFrame) -> pd.DataFrame:
    missing = REQUIRED_COLUMNS - set(df.columns)
    if missing:
        raise ValueError(f"Missing required columns: {sorted(missing)}")
    out = df.copy()
    out["Period"] = out["Period"].astype(str).str.strip()
    for col in [
        "Electricity_MWh",
        "NaturalGas_MWh",
        "Steam_MWh",
        "Production_t",
        "OperatingHours_h",
    ]:
        out[col] = pd.to_numeric(out[col], errors="coerce")
    if "Notes" not in out.columns:
        out["Notes"] = ""
    out = out[out["Period"].str.len() > 0].copy()
    if len(out) < 6:
        raise ValueError("Need at least 6 periods.")
    return out


def _sorted_periods(df: pd.DataFrame) -> list[str]:
    return sorted(df["Period"].dropna().astype(str).unique(), key=parse_period_to_datetime)


def _actions_for_display(actions: pd.DataFrame) -> pd.DataFrame:
    if len(actions) == 0:
        return pd.DataFrame(columns=ACTION_COLUMNS)
    for col in ACTION_COLUMNS:
        if col not in actions.columns:
            actions[col] = ""
    return actions[ACTION_COLUMNS]


def _validate_data_quality(df: pd.DataFrame) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    bad_periods = []
    for p in df["Period"].astype(str):
        try:
            parse_period_to_datetime(p)
        except Exception:
            bad_periods.append(p)
    if bad_periods:
        errors.append(f"Invalid Period format found: {bad_periods[:5]}. Use YYYY-MM or YYYY-Www.")
    if df["Period"].duplicated().any():
        dup_count = int(df["Period"].duplicated().sum())
        errors.append(f"Found {dup_count} duplicate period rows.")
    for col in [
        "Electricity_MWh",
        "NaturalGas_MWh",
        "Steam_MWh",
        "Production_t",
        "OperatingHours_h",
    ]:
        negatives = int((df[col] < 0).fillna(False).sum())
        if negatives > 0:
            errors.append(f"Column {col} has {negatives} negative values.")
    zero_prod = int((df["Production_t"] == 0).fillna(False).sum())
    zero_hours = int((df["OperatingHours_h"] == 0).fillna(False).sum())
    if zero_prod > 0:
        warnings.append(
            f"{zero_prod} rows have zero production; intensity may be infinite/undefined."
        )
    if zero_hours > 0:
        warnings.append(f"{zero_hours} rows have zero operating hours.")
    missing_cells = int(
        df[["Electricity_MWh", "NaturalGas_MWh", "Steam_MWh", "Production_t", "OperatingHours_h"]]
        .isna()
        .sum()
        .sum()
    )
    if missing_cells > 0:
        warnings.append(
            f"{missing_cells} missing numeric cells detected; rows may be flagged as MISSING_DATA."
        )
    return errors, warnings


def _likely_causes(period_row: pd.Series, full_df: pd.DataFrame) -> list[str]:
    causes: list[str] = []
    flags = str(period_row.get("Flags", "")).split("|")
    vals = {
        "Electricity_MWh": period_row.get("Electricity_MWh", np.nan),
        "NaturalGas_MWh": period_row.get("NaturalGas_MWh", np.nan),
        "Steam_MWh": period_row.get("Steam_MWh", np.nan),
    }
    baseline_med = full_df[["Electricity_MWh", "NaturalGas_MWh", "Steam_MWh"]].median(
        numeric_only=True
    )
    deltas = {}
    for k, v in vals.items():
        if pd.notna(v) and pd.notna(baseline_med.get(k, np.nan)) and baseline_med[k] > 0:
            deltas[k] = (v - baseline_med[k]) / baseline_med[k]
    top_utility = None
    if deltas:
        top_utility = max(deltas, key=lambda x: abs(deltas[x]))
    if "MISSING_DATA" in flags:
        causes.append("Data pipeline or meter extraction gap.")
    if "ANOMALY_RESID" in flags:
        causes.append("Operational change not explained by selected drivers.")
    if "DRIFT_HINT" in flags:
        causes.append("Sustained performance shift vs baseline behavior.")
    if top_utility == "Electricity_MWh":
        causes.append("Electrical load change: motors, compressed air, or reheating.")
    elif top_utility == "NaturalGas_MWh":
        causes.append("Fuel/process heat efficiency shift or burner tuning issue.")
    elif top_utility == "Steam_MWh":
        causes.append("Steam system losses, trap issues, or process steam demand shift.")
    while len(causes) < 3:
        causes.append("Review downtime events, product mix, and maintenance logs for this period.")
    return causes[:3]


def _log_run(metadata: dict) -> None:
    Path("app_logs").mkdir(exist_ok=True)
    out_path = Path("app_logs") / "run_metadata.jsonl"
    with out_path.open("a", encoding="utf-8") as f:
        f.write(json.dumps(metadata) + "\n")


def _persist_analysis(run_key: str, payload: dict) -> None:
    st.session_state["analysis_key"] = run_key
    st.session_state["analysis_payload"] = payload


def _set_dark_matplotlib_style() -> None:
    plt.style.use("dark_background")
    plt.rcParams["figure.facecolor"] = "#0E1117"
    plt.rcParams["axes.facecolor"] = "#0E1117"
    plt.rcParams["savefig.facecolor"] = "#0E1117"


def _plot_kpi_trends_dark(results: pd.DataFrame) -> None:
    trend = results.copy()
    trend["Date"] = pd.to_datetime(trend["Date"])
    trend = trend.sort_values("Date")

    fig1, ax1 = plt.subplots(figsize=(10, 3.8))
    ax1.plot(
        trend["Date"], trend["Intensity_kWh_per_t"], marker="o", linewidth=1.2, label="Intensity"
    )
    ax1.plot(trend["Date"], trend["Intensity_roll"], linewidth=2.0, label="Rolling")
    ax1.set_title("Energy Intensity (kWh/t)")
    ax1.set_ylabel("kWh/t")
    ax1.grid(alpha=0.2)
    ax1.legend()
    st.pyplot(fig1, clear_figure=True)

    fig2, ax2 = plt.subplots(figsize=(10, 3.8))
    ax2.plot(
        trend["Date"],
        trend["NormalizedEnergy_MWh"],
        marker="o",
        linewidth=1.2,
        label="Normalized energy",
    )
    ax2.plot(trend["Date"], trend["NormEnergy_roll"], linewidth=2.0, label="Rolling")
    ax2.set_title("Normalized Energy (MWh)")
    ax2.set_ylabel("MWh")
    ax2.grid(alpha=0.2)
    ax2.legend()
    st.pyplot(fig2, clear_figure=True)

    fig3, ax3 = plt.subplots(figsize=(10, 3.8))
    ax3.stackplot(
        trend["Date"],
        trend["Electricity_MWh"].fillna(0),
        trend["NaturalGas_MWh"].fillna(0),
        trend["Steam_MWh"].fillna(0),
        labels=["Electricity", "Natural gas", "Steam"],
        alpha=0.85,
    )
    ax3.set_title("Energy by Source (MWh)")
    ax3.set_ylabel("MWh")
    ax3.grid(alpha=0.2)
    ax3.legend(loc="upper left")
    st.pyplot(fig3, clear_figure=True)


def _render_advanced_insights(results: pd.DataFrame, rolling_window: int) -> None:
    st.subheader("Advanced Insights")
    tabs = st.tabs(
        [
            "Model Fit",
            "Residual Behavior",
            "Driver Relationships",
            "Energy Structure",
            "Seasonality",
        ]
    )

    work = results.copy()
    work["Date"] = pd.to_datetime(work["Date"])
    work = work.sort_values("Date")
    work["Month"] = work["Date"].dt.month

    with tabs[0]:
        fit_df = work.dropna(subset=["ExpectedEnergy_MWh", "TotalEnergy_MWh"])
        c1, c2 = st.columns(2)
        if len(fit_df) >= 2:
            ss_res = float(
                np.square(fit_df["TotalEnergy_MWh"] - fit_df["ExpectedEnergy_MWh"]).sum()
            )
            ss_tot = float(
                np.square(fit_df["TotalEnergy_MWh"] - fit_df["TotalEnergy_MWh"].mean()).sum()
            )
            r2 = 1.0 - (ss_res / ss_tot) if ss_tot > 1e-12 else np.nan
            mae = float((fit_df["TotalEnergy_MWh"] - fit_df["ExpectedEnergy_MWh"]).abs().mean())
            c1.metric("Model R²", f"{r2:.3f}" if np.isfinite(r2) else "-")
            c2.metric("Residual MAE (MWh)", f"{mae:.1f}")

            fig, ax = plt.subplots(figsize=(6.5, 4))
            ax.scatter(fit_df["ExpectedEnergy_MWh"], fit_df["TotalEnergy_MWh"], alpha=0.8)
            lo = min(fit_df["ExpectedEnergy_MWh"].min(), fit_df["TotalEnergy_MWh"].min())
            hi = max(fit_df["ExpectedEnergy_MWh"].max(), fit_df["TotalEnergy_MWh"].max())
            ax.plot([lo, hi], [lo, hi], linestyle="--", linewidth=1.2)
            ax.set_xlabel("Expected Energy (MWh)")
            ax.set_ylabel("Actual Total Energy (MWh)")
            ax.set_title("Actual vs Expected Energy")
            ax.grid(alpha=0.25)
            st.pyplot(fig, clear_figure=True)
        else:
            st.info("Not enough valid rows for model fit diagnostics.")

    with tabs[1]:
        resid_df = work.dropna(subset=["Residual_MWh"]).copy()
        if len(resid_df) >= 2:
            s = float(resid_df["Residual_MWh"].std(ddof=1))
            fig, ax = plt.subplots(figsize=(8.0, 4))
            ax.plot(resid_df["Date"], resid_df["Residual_MWh"], marker="o", linewidth=1.2)
            ax.axhline(0.0, linewidth=1.0)
            if np.isfinite(s) and s > 0:
                for z, style in [(2, "--"), (3, ":")]:
                    ax.axhline(z * s, linestyle=style, linewidth=1.0)
                    ax.axhline(-z * s, linestyle=style, linewidth=1.0)
            ax.set_title("Residual Control View")
            ax.set_ylabel("Residual (MWh)")
            ax.grid(alpha=0.25)
            st.pyplot(fig, clear_figure=True)

            fig2, ax2 = plt.subplots(figsize=(8.0, 3.6))
            ax2.hist(resid_df["Residual_MWh"], bins=min(12, max(6, len(resid_df) // 2)), alpha=0.85)
            ax2.set_title("Residual Distribution")
            ax2.set_xlabel("Residual (MWh)")
            ax2.set_ylabel("Frequency")
            ax2.grid(alpha=0.2)
            st.pyplot(fig2, clear_figure=True)
        else:
            st.info("Not enough residual values for distribution/control plots.")

    with tabs[2]:
        dep = work.dropna(subset=["TotalEnergy_MWh"])
        c1, c2 = st.columns(2)
        with c1:
            p = dep.dropna(subset=["Production_t"])
            if len(p) >= 2:
                fig, ax = plt.subplots(figsize=(6.0, 4))
                ax.scatter(p["Production_t"], p["TotalEnergy_MWh"], alpha=0.75)
                m, b = np.polyfit(p["Production_t"], p["TotalEnergy_MWh"], deg=1)
                x = np.linspace(p["Production_t"].min(), p["Production_t"].max(), 50)
                ax.plot(x, m * x + b, linewidth=1.2)
                ax.set_title("Energy vs Production")
                ax.set_xlabel("Production (t)")
                ax.set_ylabel("Total Energy (MWh)")
                ax.grid(alpha=0.2)
                st.pyplot(fig, clear_figure=True)
        with c2:
            h = dep.dropna(subset=["OperatingHours_h"])
            if len(h) >= 2:
                fig, ax = plt.subplots(figsize=(6.0, 4))
                ax.scatter(h["OperatingHours_h"], h["TotalEnergy_MWh"], alpha=0.75)
                m, b = np.polyfit(h["OperatingHours_h"], h["TotalEnergy_MWh"], deg=1)
                x = np.linspace(h["OperatingHours_h"].min(), h["OperatingHours_h"].max(), 50)
                ax.plot(x, m * x + b, linewidth=1.2)
                ax.set_title("Energy vs Operating Hours")
                ax.set_xlabel("Operating Hours (h)")
                ax.set_ylabel("Total Energy (MWh)")
                ax.grid(alpha=0.2)
                st.pyplot(fig, clear_figure=True)

    with tabs[3]:
        mix = work[["Date", "Electricity_MWh", "NaturalGas_MWh", "Steam_MWh"]].copy()
        total = mix[["Electricity_MWh", "NaturalGas_MWh", "Steam_MWh"]].sum(axis=1, min_count=1)
        for col in ["Electricity_MWh", "NaturalGas_MWh", "Steam_MWh"]:
            mix[f"{col}_share"] = np.where(total > 0, mix[col] / total, np.nan)
        share_df = mix.set_index("Date")[
            ["Electricity_MWh_share", "NaturalGas_MWh_share", "Steam_MWh_share"]
        ]
        fig3, ax3 = plt.subplots(figsize=(10, 3.8))
        ax3.stackplot(
            share_df.index,
            share_df["Electricity_MWh_share"].fillna(0),
            share_df["NaturalGas_MWh_share"].fillna(0),
            share_df["Steam_MWh_share"].fillna(0),
            labels=["Electricity Share", "Gas Share", "Steam Share"],
            alpha=0.9,
        )
        ax3.set_title("Utility Share Structure")
        ax3.set_ylabel("Share")
        ax3.set_ylim(0, 1)
        ax3.grid(alpha=0.2)
        ax3.legend(loc="upper left")
        st.pyplot(fig3, clear_figure=True)

        corr_df = work.set_index("Date")[
            ["TotalEnergy_MWh", "Production_t", "OperatingHours_h"]
        ].copy()
        w = max(2, int(rolling_window))
        corr_out = pd.DataFrame(index=corr_df.index)
        corr_out["Corr(Energy, Production)"] = (
            corr_df["TotalEnergy_MWh"].rolling(w).corr(corr_df["Production_t"])
        )
        corr_out["Corr(Energy, Hours)"] = (
            corr_df["TotalEnergy_MWh"].rolling(w).corr(corr_df["OperatingHours_h"])
        )
        fig4, ax4 = plt.subplots(figsize=(10, 3.8))
        ax4.plot(
            corr_out.index,
            corr_out["Corr(Energy, Production)"],
            linewidth=1.8,
            label="Energy-Production",
        )
        ax4.plot(
            corr_out.index, corr_out["Corr(Energy, Hours)"], linewidth=1.8, label="Energy-Hours"
        )
        ax4.axhline(0.0, linewidth=1.0)
        ax4.set_title(f"Rolling Correlation (window={w})")
        ax4.set_ylabel("Correlation")
        ax4.set_ylim(-1.05, 1.05)
        ax4.grid(alpha=0.2)
        ax4.legend()
        st.pyplot(fig4, clear_figure=True)

    with tabs[4]:
        season = work.dropna(subset=["TotalEnergy_MWh"]).copy()
        if len(season) >= 6:
            fig, ax = plt.subplots(figsize=(8.0, 4))
            season.boxplot(column="TotalEnergy_MWh", by="Month", ax=ax)
            ax.set_title("Monthly Distribution of Total Energy")
            ax.set_xlabel("Month")
            ax.set_ylabel("Total Energy (MWh)")
            fig.suptitle("")
            ax.grid(alpha=0.2)
            st.pyplot(fig, clear_figure=True)
        else:
            st.info("Need more rows to show seasonality distribution.")


def render() -> None:
    st.set_page_config(page_title="Energy KPI Toolkit", layout="wide")
    _set_dark_matplotlib_style()
    _auth_gate()
    st.title("Industrial Energy KPI Toolkit")
    st.caption(
        "Step 1: Upload data. Step 2: Configure baseline and thresholds. Step 3: Run analysis and export."
    )

    with st.sidebar:
        st.header("Step 1: Input")
        source = st.radio(
            "Data source", options=["Demo CSV", "Upload CSV", "Upload Excel workbook"], index=0
        )
        uploaded_csv = None
        uploaded_xlsx = None
        demo_choice = "Original demo"
        demo_map = {
            "Original demo": "demo_industrial_energy_data.csv",
            "Realistic demo (v2)": "demo_industrial_energy_data_v2.csv",
        }
        if source == "Demo CSV":
            demo_choice = st.selectbox("Demo dataset", options=list(demo_map.keys()), index=0)
        if source == "Upload CSV":
            uploaded_csv = st.file_uploader("CSV file", type=["csv"])
        elif source == "Upload Excel workbook":
            uploaded_xlsx = st.file_uploader("Excel workbook", type=["xlsx"])

    wb = None
    params_from_excel = None
    workbook_bytes = None
    try:
        if source == "Upload Excel workbook":
            if uploaded_xlsx is None:
                st.info("Upload an `.xlsx` workbook to continue.")
                st.stop()
            workbook_bytes = uploaded_xlsx.getvalue()
            wb = load_workbook(filename=BytesIO(workbook_bytes))
            params_from_excel = read_params(wb)
            df = _prepare_data(read_data_from_workbook(wb))
        elif source == "Upload CSV":
            if uploaded_csv is None:
                st.info("Upload a `.csv` file to continue.")
                st.stop()
            df = _prepare_data(_load_csv(uploaded_csv))
        else:
            df = _prepare_data(_load_csv(None, demo_map[demo_choice]))
    except Exception as exc:
        st.error(f"Could not load input data: {exc}")
        st.stop()

    errors, warnings = _validate_data_quality(df)
    for msg in warnings:
        st.warning(msg)
    if errors:
        for msg in errors:
            st.error(msg)
        st.stop()

    periods = _sorted_periods(df)
    default_start = (
        params_from_excel.baseline_start
        if (params_from_excel and params_from_excel.baseline_start in periods)
        else periods[0]
    )
    default_end = (
        params_from_excel.baseline_end
        if (params_from_excel and params_from_excel.baseline_end in periods)
        else periods[min(len(periods) - 1, 5)]
    )
    default_plant = params_from_excel.plant_name if params_from_excel else "Demo Plant"
    profile_name = "Balanced (Recommended)"

    with st.sidebar:
        st.header("Step 2: Configuration")
        plant_name = st.text_input(
            "Plant / Site name", value=default_plant, help="Used in report titles and exports."
        )
        profile_name = st.selectbox(
            "Preset profile",
            options=list(PROFILE_DEFAULTS.keys()),
            index=0,
            help="Preset threshold/model settings.",
        )
        preset = PROFILE_DEFAULTS[profile_name]
        baseline_start = st.selectbox(
            "Baseline start",
            options=periods,
            index=periods.index(default_start),
            help="Start period for model fitting.",
        )
        baseline_end = st.selectbox(
            "Baseline end",
            options=periods,
            index=periods.index(default_end),
            help="End period for model fitting.",
        )
        driver_opts = ["Both", "Production", "Hours"]
        default_driver = (
            params_from_excel.driver_model
            if (params_from_excel and params_from_excel.driver_model in driver_opts)
            else preset["driver_model"]
        )
        driver_model = st.selectbox(
            "Driver model",
            options=driver_opts,
            index=driver_opts.index(default_driver),
            help="Choose which drivers explain expected energy.",
        )
        rolling_window = st.slider(
            "Rolling window",
            min_value=1,
            max_value=12,
            value=max(1, min(12, int(preset["rolling_window"]))),
            step=1,
            help="Smooths KPI trend charts.",
        )
        anomaly_z = st.slider(
            "Anomaly z-threshold",
            min_value=1.0,
            max_value=4.0,
            value=float(preset["anomaly_z"]),
            step=0.1,
            help="Lower = more anomalies flagged.",
        )
        drift_window = st.slider(
            "Drift window",
            min_value=2,
            max_value=12,
            value=max(2, min(12, int(preset["drift_window"]))),
            step=1,
            help="Window size for drift detection.",
        )
        drift_threshold_z = st.slider(
            "Drift threshold (z)",
            min_value=0.5,
            max_value=4.0,
            value=float(preset["drift_threshold_z"]),
            step=0.1,
            help="Lower = more drift hints flagged.",
        )
        run_now = st.button("Run Analysis", type="primary")

    if parse_period_to_datetime(baseline_end) < parse_period_to_datetime(baseline_start):
        st.error("Baseline end must be the same as or after baseline start.")
        st.stop()

    params = Params(
        plant_name=plant_name,
        baseline_start=baseline_start,
        baseline_end=baseline_end,
        driver_model=driver_model,
        rolling_window=rolling_window,
        anomaly_z=anomaly_z,
        drift_window=drift_window,
        drift_threshold_z=drift_threshold_z,
    )
    input_signature = df.to_csv(index=False).encode("utf-8")
    run_key = hashlib.sha256(
        input_signature
        + f"{source}|{plant_name}|{baseline_start}|{baseline_end}|{driver_model}|{rolling_window}|{anomaly_z}|{drift_window}|{drift_threshold_z}".encode(
            "utf-8"
        )
    ).hexdigest()

    if run_now:
        try:
            results, actions, meta = compute_kpis(df, params)
        except Exception as exc:
            st.error(f"KPI calculation failed: {exc}")
            st.stop()
        payload = {
            "results": results,
            "actions": _actions_for_display(actions.copy()),
            "meta": meta,
            "params": params,
        }
        _persist_analysis(run_key, payload)
        _log_run(
            {
                "timestamp_utc": dt.datetime.utcnow().isoformat(timespec="seconds"),
                "source": source,
                "profile": profile_name,
                "plant_name": plant_name,
                "baseline_start": baseline_start,
                "baseline_end": baseline_end,
                "driver_model": driver_model,
                "rolling_window": rolling_window,
                "anomaly_z": anomaly_z,
                "drift_window": drift_window,
                "drift_threshold_z": drift_threshold_z,
                "rows": int(len(df)),
                "run_key": run_key[:12],
            }
        )

    if st.session_state.get("analysis_key") != run_key:
        st.info("Set parameters and click `Run Analysis`.")
        st.subheader("Loaded Data")
        st.dataframe(df, width="stretch", hide_index=True)
        st.stop()

    payload = st.session_state["analysis_payload"]
    results = payload["results"]
    actions_view = payload["actions"]
    meta = payload["meta"]

    st.header("Step 3: Review Results")
    latest = results.dropna(subset=["TotalEnergy_MWh"]).tail(1).iloc[0]
    open_issues = int((actions_view["Status"] == "Open").sum()) if len(actions_view) else 0
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Latest Total Energy (MWh)", f"{latest['TotalEnergy_MWh']:.1f}")
    c2.metric(
        "Latest Intensity (kWh/t)",
        f"{latest['Intensity_kWh_per_t']:.1f}" if pd.notna(latest["Intensity_kWh_per_t"]) else "-",
    )
    c3.metric(
        "Latest Normalized Energy (MWh)",
        (
            f"{latest['NormalizedEnergy_MWh']:.1f}"
            if pd.notna(latest["NormalizedEnergy_MWh"])
            else "-"
        ),
    )
    c4.metric("Open Issues", str(open_issues))

    baseline_df = results[
        (results["Date"] >= parse_period_to_datetime(baseline_start))
        & (results["Date"] <= parse_period_to_datetime(baseline_end))
    ]
    base_mean = float(baseline_df["NormalizedEnergy_MWh"].mean()) if len(baseline_df) else np.nan
    latest_norm = (
        float(latest["NormalizedEnergy_MWh"])
        if pd.notna(latest["NormalizedEnergy_MWh"])
        else np.nan
    )
    if np.isfinite(base_mean) and np.isfinite(latest_norm):
        delta_pct = (
            ((latest_norm - base_mean) / base_mean) * 100 if abs(base_mean) > 1e-9 else np.nan
        )
        direction = "above" if delta_pct >= 0 else "below"
        st.info(
            f"Interpretation: latest normalized energy is {abs(delta_pct):.1f}% {direction} the baseline average."
        )

    st.subheader("Data Overview")
    st.dataframe(df, width="stretch", hide_index=True)

    st.subheader("KPI Trends")
    _plot_kpi_trends_dark(results)
    _render_advanced_insights(results, rolling_window=params.rolling_window)

    st.subheader("Flag Drill-Down")
    flagged = results[results["Flags"].astype(str).str.len() > 0].copy()
    if len(flagged) == 0:
        st.info("No flagged periods in this run.")
    else:
        selected = st.selectbox("Select flagged period", options=flagged["Period"].tolist())
        row = flagged[flagged["Period"] == selected].iloc[0]
        d1, d2 = st.columns(2)
        d1.write(
            {
                "Period": row["Period"],
                "Flags": row["Flags"],
                "TotalEnergy_MWh": (
                    float(row["TotalEnergy_MWh"]) if pd.notna(row["TotalEnergy_MWh"]) else None
                ),
                "Residual_MWh": (
                    float(row["Residual_MWh"]) if pd.notna(row["Residual_MWh"]) else None
                ),
                "z_resid": float(row["z_resid"]) if pd.notna(row["z_resid"]) else None,
            }
        )
        d2.markdown("Likely causes:")
        for cause in _likely_causes(row, results):
            d2.write(f"- {cause}")

    st.subheader("Results Table")
    st.dataframe(results, width="stretch", hide_index=True)

    st.subheader("Actions Log")
    if len(actions_view) == 0:
        st.info("No issues were flagged for the selected settings.")
        edited_actions = actions_view
    else:
        edited_actions = st.data_editor(
            actions_view,
            width="stretch",
            hide_index=True,
            num_rows="dynamic",
            column_config={
                "Severity": st.column_config.SelectboxColumn(options=["Low", "Med", "High"]),
                "Status": st.column_config.SelectboxColumn(
                    options=["Open", "In Progress", "Closed"]
                ),
            },
        )

    st.subheader("Exports")
    st.download_button(
        "Download results.csv",
        data=results.to_csv(index=False).encode("utf-8"),
        file_name="results.csv",
        mime="text/csv",
    )
    st.download_button(
        "Download actions.csv",
        data=edited_actions.to_csv(index=False).encode("utf-8"),
        file_name="actions.csv",
        mime="text/csv",
    )

    if wb is not None and workbook_bytes is not None:
        wb_for_download = load_workbook(filename=BytesIO(workbook_bytes))
        write_results_to_workbook(wb_for_download, results, edited_actions)
        out_bytes = BytesIO()
        wb_for_download.save(out_bytes)
        st.download_button(
            "Download updated_workbook.xlsx",
            data=out_bytes.getvalue(),
            file_name="updated_workbook.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    with tempfile.TemporaryDirectory() as tmpdir:
        plot_dir = Path(tmpdir) / "plots"
        img_paths = make_plots(results, params, str(plot_dir))
        pdf_path = Path(tmpdir) / "management_report.pdf"
        make_pdf_report(results, edited_actions, params, img_paths, str(pdf_path))
        st.download_button(
            "Download management_report.pdf",
            data=pdf_path.read_bytes(),
            file_name="management_report.pdf",
            mime="application/pdf",
        )

    with st.expander("Model details"):
        st.write(f"Baseline rows used: {meta['baseline_rows']}")
        st.write(f"Model features: {meta['feat_names']}")
        st.write(f"Coefficients: {np.round(meta['beta'], 4).tolist()}")
        st.write(f"Residual std (baseline): {meta['resid_std']:.4f}")
        st.write("Run metadata log: `app_logs/run_metadata.jsonl`")

    with st.expander("Scientific and Guideline References Used"):
        st.write(
            "Methods in this demo align with standard EnPI/baseline and M&V practice "
            "(regression normalization, whole-facility tracking, and residual/change diagnostics)."
        )
        for label, url in REFERENCE_LINKS:
            st.markdown(f"- [{label}]({url})")

    with st.expander("Paper-to-Plot Mapping"):
        st.write("How each advanced diagnostic maps to published methods/guidelines.")
        st.dataframe(pd.DataFrame(PAPER_PLOT_MAPPING), width="stretch", hide_index=True)


if __name__ == "__main__":
    render()
